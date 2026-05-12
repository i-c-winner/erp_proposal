# Copyright (c) 2026, Belousov Dmitriy and contributors
# For license information, please see license.txt

import base64
from io import BytesIO

import frappe
from frappe.model.document import Document
from frappe.model.naming import make_autoname


STATUS_NOTIFICATIONS = {
    "Budget Pending": {
        "role": "Sales Manager",
        "subject": "Подготовить бюджет к коммерческому предложению {name}",
    },
    "Budget Pending Approval": {
        "role": "Sales Master Manager",
        "subject": "Согласовать бюджет {name}",
    },
    "Budget Approved": {
        "role": "Sales Manager",
        "subject": "Подготовить коммерческое предложение {name}",
    },
    "Proposal Pending": {
        "role": "Sales Master Manager",
        "subject": "Согласовать коммерческое предложение {name}",
    },
    "Final Approval": {
        "role": "Sales Manager",
        "subject": "Финальное утверждение прошло {name}",
    },
}

UNREAD_NOTIFICATIONS_CHANGED_EVENT = "proposal_unread_notifications_changed"


class CommercialProposal(Document):
    def before_insert(self):
        if self.status == "Draft":
            self.status = "Budget Pending"

    def after_insert(self):
        self.attach_budget_template()
        self.notify_for_status()

    def attach_budget_template(self):
        """Store default budget template content in DB when a new document is created."""
        # Use DB query — self may not yet have budget_content attribute if meta cache is stale
        existing = frappe.db.get_value("Commercial Proposal", self.name, "budget_content")
        if existing:
            return
        try:
            content_b64 = _build_template_b64()
            safe_name = self.name.replace("/", "-")
            frappe.db.set_value(
                "Commercial Proposal",
                self.name,
                {
                    "budget": f"budget_{safe_name}.xlsx",
                    "budget_content": content_b64,
                },
                update_modified=False,
            )
        except Exception:
            frappe.log_error(frappe.get_traceback(), "Budget Template Attach Error")

    def on_update(self):
        if self.has_value_changed("status"):
            self.notify_for_status()

    def autoname(self):
        prefix = (self.prefix or "PREFIX").strip().upper()
        work_type = (self.work_type or "TYPE").strip().upper()
        self.name = make_autoname(f"{prefix}/{work_type}/.YYYY./.MM./.#####")

    def validate(self):
        self.validate_budget_file()
        self.calculate_totals()

    def validate_budget_file(self):
        if self.budget:
            fname = self.budget.lower()
            if not (fname.endswith(".xlsx") or fname.endswith(".xls")):
                frappe.throw(
                    frappe._("Only Excel files (.xlsx, .xls) are allowed in the Budget field.")
                )

    def calculate_totals(self):
        total = 0
        self.currency = self.currency or "UZS"
        for row in self.budget_lines:
            row.total_amount = (row.qty or 0) * (row.rate or 0)
            total += row.total_amount
        self.total_amount = total

    def notify_for_status(self):
        notification = STATUS_NOTIFICATIONS.get(self.status)
        if not notification:
            return
        users = get_users_with_role(notification["role"])
        if not users:
            return
        subject = notification["subject"].format(name=self.name)
        for user in users:
            create_notification_log(self, user, subject)


# ─── Budget helpers ──────────────────────────────────────────────────────────

def _build_template_b64():
    """Read the budget Excel template from the app's templates directory and return as base64 string."""
    template_path = frappe.get_app_path("proposal", "templates", "budget_template.xlsx")
    with open(template_path, "rb") as f:
        return base64.b64encode(f.read()).decode("ascii")


def _read_cells_from_b64(content_b64):
    """Return {total_amount, currency} from cells B2 and C2 of the Excel content."""
    import openpyxl

    result = {}
    try:
        data = base64.b64decode(content_b64)
        wb = openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
        ws = wb.active

        raw_total = ws["B2"].value
        raw_currency = ws["C2"].value
        wb.close()

        if raw_total is not None:
            try:
                result["total_amount"] = float(raw_total)
            except (ValueError, TypeError):
                pass
        if raw_currency is not None:
            result["currency"] = str(raw_currency).strip()
    except Exception:
        frappe.log_error(frappe.get_traceback(), "Budget cell sync error")

    return result


# ─── Whitelisted API ─────────────────────────────────────────────────────────

@frappe.whitelist()
def save_budget_file(docname, file_content_b64, file_name):
    """Save base64 Excel content from the in-app editor directly to the database."""
    doc = frappe.get_doc("Commercial Proposal", docname)
    frappe.has_permission("Commercial Proposal", "write", doc=doc, throw=True)

    file_name = frappe.utils.sanitize_html(file_name or "budget.xlsx")
    if not (file_name.lower().endswith(".xlsx") or file_name.lower().endswith(".xls")):
        frappe.throw(frappe._("Only Excel files (.xlsx, .xls) are allowed."))

    cell_values = _read_cells_from_b64(file_content_b64)
    update = {"budget": file_name, "budget_content": file_content_b64}
    update.update(cell_values)

    frappe.db.set_value("Commercial Proposal", docname, update, update_modified=False)
    return file_name


@frappe.whitelist()
def get_budget_file_content(docname):
    """Return base64 Excel content stored in the document."""
    doc = frappe.get_doc("Commercial Proposal", docname)
    frappe.has_permission("Commercial Proposal", "read", doc=doc, throw=True)

    row = frappe.db.get_value(
        "Commercial Proposal", docname, ["budget_content", "budget"], as_dict=True
    )
    if not row or not row.budget_content:
        frappe.throw(frappe._("No budget file attached."))

    return {
        "content_b64": row.budget_content,
        "file_name": row.budget or "budget.xlsx",
    }



# ─── Notification helpers ────────────────────────────────────────────────────

def get_users_with_role(role):
    users = frappe.get_all(
        "Has Role",
        filters={"role": role, "parenttype": "User"},
        pluck="parent",
    )
    if not users:
        return []
    return frappe.get_all(
        "User",
        filters={"name": ("in", users), "enabled": 1},
        pluck="name",
    )


def create_notification_log(doc, user, subject):
    if notification_exists(doc, user, subject):
        return
    notification = frappe.new_doc("Notification Log")
    notification.update(
        {
            "subject": subject,
            "for_user": user,
            "from_user": frappe.session.user,
            "type": "Alert",
            "document_type": doc.doctype,
            "document_name": doc.name,
            "email_content": subject,
        }
    )
    notification.insert(ignore_permissions=True)
    frappe.publish_realtime(UNREAD_NOTIFICATIONS_CHANGED_EVENT, user=user, after_commit=True)
    frappe.publish_realtime(
        "proposal_notification",
        {"document_type": doc.doctype, "document_name": doc.name, "subject": subject},
        user=user,
        after_commit=True,
    )


def notification_exists(doc, user, subject):
    return frappe.db.exists(
        "Notification Log",
        {
            "for_user": user,
            "subject": subject,
            "document_type": doc.doctype,
            "document_name": doc.name,
        },
    )


@frappe.whitelist()
def get_unread_notification_count():
    if frappe.session.user == "Guest":
        return 0
    return frappe.db.count(
        "Notification Log",
        filters={"for_user": frappe.session.user, "read": 0},
    )


@frappe.whitelist()
def get_notification_logs(limit=20):
    if frappe.session.user == "Guest":
        return {"notification_logs": [], "user_info": {}}
    notification_logs = frappe.get_all(
        "Notification Log",
        filters={"for_user": frappe.session.user},
        fields=["*"],
        order_by="creation desc",
        limit=int(limit or 20),
        ignore_permissions=True,
    )
    user_info = frappe._dict()
    for user in {log.from_user for log in notification_logs if log.from_user}:
        frappe.utils.add_user_info(user, user_info)
    return {"notification_logs": notification_logs, "user_info": user_info}


@frappe.whitelist()
def get_notification_debug():
    if frappe.session.user == "Guest":
        return {"user": "Guest", "unread_count": 0, "recent": []}
    return {
        "user": frappe.session.user,
        "unread_count": get_unread_notification_count(),
        "recent": frappe.get_all(
            "Notification Log",
            filters={"for_user": frappe.session.user},
            fields=["name", "for_user", "read", "subject", "creation"],
            order_by="creation desc",
            limit=5,
            ignore_permissions=True,
        ),
    }
